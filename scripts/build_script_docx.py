import json, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "한컴돋움"   # HWP 원본 서식: 12pt, 장평 83%, 자간 -5%, 줄간격 145%
BASE = 12
LINE = Pt(BASE * 1.45)
CUE = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x80, 0x80, 0x80)
TOKEN = re.compile(r"(\*\*[^*]+\*\*|\([^)]*\)|<[^>]*>|\[[^\]]*\])")

def set_font(run, size=BASE, bold=False, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    for tag, val in (("w:w", "83"), ("w:spacing", str(round(-0.05 * size * 20)))):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); rpr.append(el)
        el.set(qn("w:val"), val)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color

def to_min(t):
    """'2분30초' / '2.5분' / '90초' -> 분(소수). 없으면 None"""
    if not t:
        return None
    m = re.search(r"(\d+(?:\.\d+)?)\s*분", t); s = re.search(r"(\d+)\s*초", t)
    if not (m or s):
        return None
    return round((float(m.group(1)) if m else 0) + (int(s.group(1)) / 60 if s else 0), 2)

def add_field(p, instr, cached, size=12, bold=True, color=None):
    def r(child):
        run = p.add_run(); set_font(run, size=size, bold=bold, color=color)
        run._element.append(child)
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r(it)
    sp = OxmlElement("w:fldChar"); sp.set(qn("w:fldCharType"), "separate"); r(sp)
    cr = p.add_run(cached); set_font(cr, size=size, bold=bold, color=color)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r(e)

def add_bookmarked(p, text, name, bid, **font):
    bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), str(bid)); bs.set(qn("w:name"), name)
    be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), str(bid))
    p._p.append(bs)
    run = p.add_run(text); set_font(run, **font)
    p._p.append(be)

def shade(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)

def fmt_para(p):
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = LINE  # 고정 줄간격 = 12pt x 145%

def write_line(cell, text, first, size=BASE):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    fmt_para(p)
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            set_font(p.add_run(part[2:-2]), size=size, bold=True)
            continue
        r = p.add_run(part)
        if part.startswith("[") and part.endswith("]"):
            set_font(r, size=size, bold=True); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part[0] in "(<":
            set_font(r, size=size - 1, bold=True, color=CUE)
        else:
            set_font(r, size=size)

def build(spec, out):
    m = spec["meta"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.0)
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(t.add_run("결혼식순 WEDDING PROGRAM"), size=16, bold=True)
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = f"신랑 {m['groom']} · 신부 {m['bride']}  /  {m['date']}  /  {m['time']}  /  {m['venue']}"
    set_font(h.add_run(info), size=10)
    if m.get("note"):
        n = doc.add_paragraph(); n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(n.add_run(m["note"]), size=9.5, color=CUE)

    # 순서별 시간 -> 책갈피 이름 (T1, T2 …)
    mins = [to_min(r.get("time")) for r in spec["rows"]]
    names = {i: f"T{i + 1}" for i, v in enumerate(mins) if v is not None}
    total = sum(v for v in mins if v is not None)
    expr = "+".join(names.values()) or "0"

    tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(tp.add_run("총 소요시간  "), size=12, bold=True)
    add_field(tp, f'=INT({expr}) \\# "0"', str(int(total)))
    set_font(tp.add_run("분 "), size=12, bold=True)
    add_field(tp, f'=MOD(({expr})*60,60) \\# "0"', str(round(total * 60) % 60))
    set_font(tp.add_run("초"), size=12, bold=True)
    hp = doc.add_paragraph(); hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(hp.add_run("※ 진행순서 칸의 분 숫자를 고친 뒤 Ctrl+A → F9 를 누르면 총 소요시간이 다시 계산됩니다 (2분30초 = 2.5)"), size=8.5, color=GRAY)

    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Cm(2.2), Cm(14.13), Cm(2.55))   # HWP 원본 표: 22 / 141.3 / 25.5 mm
    for i, label in enumerate(("진행순서", "진행대본", "비고")):
        c = tbl.rows[0].cells[i]; c.width = widths[i]; shade(c, "D9E2F3")
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; fmt_para(p)
        set_font(p.add_run(label), bold=True)
    tbl.autofit = False
    for i, w in enumerate(widths):
        tbl.columns[i].width = w
    for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))

    for idx, row in enumerate(spec["rows"]):
        cells = tbl.add_row().cells
        for c, w in zip(cells, widths):
            c.width = w
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cells[0], "F2F2F2")
        p = cells[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; fmt_para(p)
        set_font(p.add_run(row["step"]), bold=True)
        if idx in names:
            p2 = cells[0].add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; fmt_para(p2)
            set_font(p2.add_run("("), size=11, color=CUE)
            add_bookmarked(p2, "%g" % mins[idx], names[idx], idx + 1, size=11, bold=True, color=CUE)
            set_font(p2.add_run("분)"), size=11, color=CUE)
        for i, line in enumerate(row["lines"]):
            write_line(cells[1], line, i == 0)
        cells[2].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for i, memo in enumerate(row.get("memo", [])):
            write_line(cells[2], memo, i == 0, size=10)
        for c in cells[1:]:
            fmt_para(c.paragraphs[0])

    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
    doc.settings.element.append(uf)
    doc.save(out)

if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f), sys.argv[2])
    print("saved", sys.argv[2])
